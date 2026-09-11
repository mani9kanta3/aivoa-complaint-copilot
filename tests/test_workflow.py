from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from pydantic import ValidationError

from backend import ai
from backend.documents import extract_document
from backend.schemas import ComplaintData, Extraction, RiskAssessment


def state(data=None, document_text=''):
    return {
        'message': 'Change batch to BMX24602 and quantity to 48 capsules.',
        'document_text': document_text,
        'source': 'Chat',
        'data': data or ComplaintData().model_dump(),
        'is_edit': bool(data),
        'risk': {},
        'reply': '',
        'changed_fields': [],
        'missing_fields': [],
        'steps': [],
    }


def fake_response(instructions, payload, schema):
    if schema is Extraction:
        return Extraction(changes={'batch_number': 'BMX24602', 'quantity_affected': '48 capsules'}, reply='Batch and quantity updated.')
    return RiskAssessment(
        severity='Major', priority='High', summary='Discolored capsules require review.',
        rationale='Appearance defect reported.', next_action='Route to QA investigation.',
        potential_root_causes=['Possible moisture exposure.'], capa_recommendations=['Investigate storage records.'],
    )


def test_edit_preserves_unrelated_fields_and_reassesses_risk(monkeypatch):
    monkeypatch.setattr(ai, 'ask_groq', fake_response)
    data = ComplaintData(product_name='Amoxicillin capsules', batch_number='OLD', customer_name='Apollo Pharmacy', description='Brown spots on capsule shells.').model_dump()
    result = ai.complaint_graph.invoke(state(data))
    assert result['data']['batch_number'] == 'BMX24602'
    assert result['data']['quantity_affected'] == '48 capsules'
    for key in ['product_name', 'customer_name', 'description', 'manufacturing_date']:
        assert result['data'][key] == data[key]
    assert result['risk']['severity'] == 'Major'
    assert result['steps'] == ['Edit complaint', 'Risk assessment', 'Completeness check']
    assert 'Customer contact' in result['missing_fields']


def test_document_uses_document_graph_node(monkeypatch):
    monkeypatch.setattr(ai, 'ask_groq', fake_response)
    result = ai.complaint_graph.invoke(state(document_text='A customer has reported discolored capsules.'))
    assert result['steps'][0] == 'Document extraction'


def test_greeting_does_not_invent_fields_or_assess_risk(monkeypatch):
    def greeting(instructions, payload, schema):
        assert schema is Extraction
        return Extraction(reply='Please share complaint details.')
    monkeypatch.setattr(ai, 'ask_groq', greeting)
    result = ai.complaint_graph.invoke(state())
    assert not any(result['data'].values())
    assert result['risk'] == {}


def test_explicit_clear_removes_only_requested_field(monkeypatch):
    monkeypatch.setattr(ai, 'ask_groq', lambda *args: Extraction(clear_fields=['customer_contact'], reply='Contact removed.'))
    data = ComplaintData(customer_contact='old@example.com', product_name='Metformin').model_dump()
    result = ai.edit_complaint(state(data))
    assert result['data']['customer_contact'] == ''
    assert result['data']['product_name'] == 'Metformin'


def test_ai_cannot_add_status_or_unknown_fields():
    with pytest.raises(ValidationError):
        Extraction(changes={'status': 'Approved'}, reply='Approved.')
    with pytest.raises(ValidationError):
        Extraction(clear_fields=['version'], reply='Removed.')


@pytest.mark.parametrize('filename,content', [('file.exe', b'bad file'), ('file.pdf', b''), ('file.txt', b'x' * (10 * 1024 * 1024 + 1)), ('file.txt', b'tiny')], ids=['unsupported', 'empty', 'too-large', 'too-short'])
def test_invalid_uploads(filename, content):
    with pytest.raises(ValueError):
        extract_document(filename, content)


def test_eml_extracts_customer_and_defect():
    path = Path('samples/amoxicillin-complaint.eml')
    result = extract_document(path.name, path.read_bytes())
    assert 'quality@apollo.example' in result
    assert 'AMX260701' in result
    assert 'brown spots' in result


def test_docx_reads_tables():
    document = Document()
    document.add_paragraph('Customer complaint about torn packaging.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Batch number'
    table.cell(0, 1).text = 'MFH260712A'
    stream = BytesIO()
    document.save(stream)
    assert 'MFH260712A' in extract_document('sample.docx', stream.getvalue())


def test_pdf_extracts_sample_fields():
    path = Path('samples/metformin-complaint.pdf')
    result = extract_document(path.name, path.read_bytes())
    assert 'MFH260712A' in result
    assert '25 kg' in result
    assert 'torn inner liner' in result
