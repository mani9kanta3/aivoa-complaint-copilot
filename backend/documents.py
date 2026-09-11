from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from backend.config import MAX_FILE_SIZE, MAX_TEXT_LENGTH


class EmailHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)


def extract_document(filename, content):
    extension = Path(filename).suffix.lower()
    if extension not in ['.pdf', '.docx', '.txt', '.eml']:
        raise ValueError('Upload a PDF, DOCX, TXT, or EML file.')
    if not content:
        raise ValueError('The uploaded file is empty.')
    if len(content) > MAX_FILE_SIZE:
        raise ValueError('The file is too large. The limit is 10 MB.')

    try:
        if extension == '.pdf':
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError('Upload a PDF without password protection.')
            if len(reader.pages) > 50:
                raise ValueError('Please upload a PDF with 50 pages or fewer.')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        elif extension == '.docx':
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError('The DOCX expands beyond the supported size.')
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(' | '.join(cell.text for cell in row.cells))
            text = '\n'.join(parts)
        elif extension == '.eml':
            email = BytesParser(policy=policy.default).parsebytes(content)
            body = email.get_body(preferencelist=('plain', 'html'))
            text = body.get_content() if body else ''
            if body and body.get_content_type() == 'text/html':
                parser = EmailHTML()
                parser.feed(text)
                text = ' '.join(parser.parts)
            text = f"From: {email.get('From', '')}\nDate: {email.get('Date', '')}\nSubject: {email.get('Subject', '')}\n\n{text}"
        else:
            text = content.decode('utf-8-sig')
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('This file could not be read. Try a text-based PDF, DOCX, TXT, or EML file.') from error

    text = text.replace('\x00', '').strip()
    if len(text) < 20:
        raise ValueError('No usable complaint text was found. Scanned PDFs need OCR; please paste the text instead.')
    if len(text) > MAX_TEXT_LENGTH:
        raise ValueError('The document contains too much text. Upload a shorter complaint, up to 24,000 characters.')
    return text
